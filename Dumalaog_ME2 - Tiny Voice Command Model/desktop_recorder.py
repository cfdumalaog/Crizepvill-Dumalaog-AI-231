"""Native Windows recorder for ME2 VCM; captures directly with PortAudio."""
from __future__ import annotations

import csv
import os
from pathlib import Path
import sys
import threading
import time
import tkinter as tk
from tkinter import messagebox, ttk

import numpy as np
from PIL import Image, ImageTk
import sounddevice as sd
from scipy.signal import resample_poly, spectrogram as scipy_spectrogram

from tinyvcm_model.config import SR
from tinyvcm_model.manifest_export import workbook_payload
from tinyvcm_model.recording_labels import (
    BACKGROUND_LABEL, RECORDER_LABELS, RECORDER_PHRASES, RECORDER_VARIATIONS,
    SILENCE_LABEL, WAKE_LABEL,
)
from tinyvcm_model.recorder_core import save_recording_data


APP_TITLE = 'ME2 - VCM on Raspberry Pi 5 — Dataset Recorder'
CONDITIONS = ('quiet-near', 'fan-near', 'quiet-far')


def project_root() -> Path:
    anchor = Path(sys.executable if getattr(sys, 'frozen', False) else __file__).resolve()
    for candidate in (anchor.parent, *anchor.parents):
        if (candidate / 'docs' / 'ground_truth_phrases').is_dir() and (candidate / 'tinyvcm_model').is_dir():
            return candidate
    raise RuntimeError('Cannot find the ME2 project root. Keep this app inside the project launchers folder.')


def excel_export(root: Path) -> Path:
    """Create the same audit workbook locally without a browser or Node runtime."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    payload = workbook_payload(root)
    book = Workbook()
    book.remove(book.active)

    def make_sheet(name, subtitle, headers, rows, widths):
        sheet = book.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.append([payload['title']])
        sheet.append([subtitle])
        sheet.append([f"Exported {payload['created_at']}; active pair {payload['active_run']}"])
        sheet.append([])
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(name='Arial', size=15, bold=True, color='202020')
        for row_number in (2, 3):
            sheet.cell(row_number, 1).font = Font(name='Arial', size=10, color='555555')
        for cell in sheet[5]:
            cell.fill = PatternFill('solid', fgColor='26374B')
            cell.font = Font(name='Arial', bold=True, color='FFFFFF')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        for row in sheet.iter_rows(min_row=6):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[5].height = 38
        sheet.freeze_panes = 'B6'
        for col, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.auto_filter.ref = f'A5:{get_column_letter(len(headers))}{sheet.max_row}'
        return sheet

    make_sheet('Command phrases', 'Canonical prompts come from the same files used by the recorder. Spoken transcripts are not inferred.',
               ['Label', 'Canonical suggested phrase', 'Training role', 'Matching reference transcripts', 'Raw saved takes', 'Unique clips in active run', 'Ground truth source'],
               payload['commands'], [33, 52, 20, 21, 16, 22, 54])
    make_sheet('Recordings', f"{payload['raw_rows']} saved rows; {payload['logged_prompt_rows']} logged prompts. Blank historical prompts mean unknown.",
               ['Row', 'Original label', 'Mapped intent', 'Current canonical phrase', 'Prompt logged at save', 'Prompt variant index', 'Prompt provenance', 'Spoken transcript status', 'Speaker ID', 'Condition', 'Active run split / status', 'Duplicate group size', 'Source ID', 'Relative WAV path', 'PCM SHA-256'],
               payload['recordings'], [7, 31, 31, 48, 48, 18, 28, 29, 16, 17, 30, 19, 38, 65, 72])
    performance = make_sheet('Intent performance',
               f"Reference accuracy {100*payload['reference_accuracy']:.2f}%; personal accuracy {100*payload['personal_accuracy']:.2f}%.",
               ['Intent', 'Exact recorder phrase', 'Reference recall', 'Reference F1', 'Reference test n', 'Personal precision', 'Personal recall', 'Personal F1', 'Personal test n', 'Recording priority', 'Live confidence gate', 'Live margin gate', 'Evaluation policy'],
               payload['performance'], [32, 50, 18, 18, 18, 18, 18, 18, 18, 24, 20, 20, 33])
    for row in performance.iter_rows(min_row=6, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = '0.00%'
    for row in performance.iter_rows(min_row=6, min_col=6, max_col=8):
        for cell in row:
            cell.number_format = '0.00%'
    make_sheet('Wake performance',
               f"{payload['snapshot_rows']}-row frozen training snapshot. Cutoff is confidence, not accuracy.",
               ['Operating point', 'Wake cutoff', 'Wake hits', 'Wake test n', 'Wake recall', 'Wake precision', 'Wake F1', 'Binary accuracy', 'Personal false accepts', 'Personal negatives n', 'Reference false accepts', 'Reference negatives n'],
               payload['wake'], [49, 19, 14, 15, 18, 18, 18, 20, 23, 22, 24, 23])
    output = root / 'data' / 'human' / 'manifest.xlsx'
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    return output


class DesktopRecorder:
    def __init__(self, root: Path):
        self.root = root
        self.window = tk.Tk()
        self.window.title(APP_TITLE)
        self.window.geometry('1120x820')
        self.window.minsize(860, 680)
        self.window.protocol('WM_DELETE_WINDOW', self.close)
        self.input_devices: list[tuple[int, str, float, str]] = []
        self.output_devices: list[tuple[int, str, float, str]] = []
        self.chunks: list[np.ndarray] = []
        self.chunk_lock = threading.Lock()
        self.stream = None
        self.recorded_audio: tuple[int, np.ndarray] | None = None
        self.recording = False
        self.record_started_at = 0.0
        self.audio_warning = None

        self.speaker = tk.StringVar(value='person-01')
        self.condition = tk.StringVar(value=CONDITIONS[0])
        self.label = tk.StringVar(value=WAKE_LABEL)
        self.phrase = tk.StringVar(value=RECORDER_PHRASES[WAKE_LABEL])
        self.input_choice = tk.StringVar()
        self.output_choice = tk.StringVar()
        self.consent = tk.BooleanVar(value=False)
        self.status = tk.StringVar(value='Ready — choose a microphone and click Start recording.')
        self.duration = tk.StringVar(value='Duration: 0.00 s')

        self._build_ui()
        self.refresh_devices()
        self.window.after(150, self._update_live_view)

    def _build_ui(self):
        style = ttk.Style(self.window)
        try:
            style.theme_use('vista')
        except tk.TclError:
            pass
        outer = ttk.Frame(self.window, padding=14)
        outer.pack(fill='both', expand=True)
        ttk.Label(outer, text=APP_TITLE, font=('Segoe UI', 17, 'bold')).pack(anchor='w')
        ttk.Label(outer, text='Offline desktop recorder — audio stays on this PC; no web server or browser microphone permission is used.',
                  wraplength=1040).pack(anchor='w', pady=(3, 12))

        metadata = ttk.LabelFrame(outer, text='Recording details', padding=10)
        metadata.pack(fill='x')
        ttk.Label(metadata, text='Anonymous speaker ID').grid(row=0, column=0, sticky='w')
        ttk.Entry(metadata, textvariable=self.speaker, width=22).grid(row=0, column=1, padx=(8, 18), sticky='w')
        ttk.Label(metadata, text='Condition').grid(row=0, column=2, sticky='w')
        ttk.Combobox(metadata, textvariable=self.condition, values=CONDITIONS, state='readonly', width=18).grid(row=0, column=3, padx=(8, 18), sticky='w')
        ttk.Label(metadata, text='Label').grid(row=0, column=4, sticky='w')
        label_box = ttk.Combobox(metadata, textvariable=self.label, values=RECORDER_LABELS, state='readonly', width=34)
        label_box.grid(row=0, column=5, padx=(8, 0), sticky='ew')
        metadata.columnconfigure(5, weight=1)
        label_box.bind('<<ComboboxSelected>>', self._label_changed)
        ttk.Label(metadata, text='Suggested phrase / variation').grid(row=1, column=0, sticky='w', pady=(12, 0))
        self.phrase_box = ttk.Combobox(metadata, textvariable=self.phrase, state='readonly')
        self.phrase_box.grid(row=1, column=1, columnspan=5, sticky='ew', padx=(8, 0), pady=(12, 0))
        self._refresh_phrases()
        ttk.Checkbutton(metadata, text='Speaker agrees to these local course-project recordings', variable=self.consent).grid(row=2, column=0, columnspan=6, sticky='w', pady=(12, 0))

        devices = ttk.LabelFrame(outer, text='Audio devices', padding=10)
        devices.pack(fill='x', pady=(10, 0))
        ttk.Label(devices, text='Microphone').grid(row=0, column=0, sticky='w')
        self.input_box = ttk.Combobox(devices, textvariable=self.input_choice, state='readonly')
        self.input_box.grid(row=0, column=1, sticky='ew', padx=8)
        ttk.Label(devices, text='Speaker for playback').grid(row=0, column=2, sticky='w')
        self.output_box = ttk.Combobox(devices, textvariable=self.output_choice, state='readonly')
        self.output_box.grid(row=0, column=3, sticky='ew', padx=8)
        ttk.Button(devices, text='↻ Refresh devices', command=self.refresh_devices).grid(row=0, column=4, sticky='e')
        devices.columnconfigure(1, weight=1)
        devices.columnconfigure(3, weight=1)

        controls = ttk.Frame(outer)
        controls.pack(fill='x', pady=(12, 6))
        self.start_button = ttk.Button(controls, text='● Start recording', command=self.start_recording)
        self.start_button.pack(side='left')
        self.stop_button = ttk.Button(controls, text='■ Stop', command=self.stop_recording, state='disabled')
        self.stop_button.pack(side='left', padx=8)
        self.play_button = ttk.Button(controls, text='▶ Play take', command=self.play_take, state='disabled')
        self.play_button.pack(side='left')
        self.save_button = ttk.Button(controls, text='Save take', command=self.save_take, state='disabled')
        self.save_button.pack(side='left', padx=8)
        ttk.Button(controls, text='Clear / record another', command=self.clear_take).pack(side='left')
        ttk.Button(controls, text='Refresh Excel manifest', command=self.export_excel).pack(side='right')

        status_frame = ttk.Frame(outer)
        status_frame.pack(fill='x', pady=(4, 4))
        self.status_label = ttk.Label(status_frame, textvariable=self.status, font=('Segoe UI', 10, 'bold'), wraplength=850)
        self.status_label.pack(side='left', fill='x', expand=True)
        ttk.Label(status_frame, textvariable=self.duration).pack(side='right')

        visualization = ttk.LabelFrame(outer, text='Live spectrogram — updates while recording', padding=6)
        visualization.pack(fill='both', expand=True, pady=(6, 8))
        self.spectrogram_label = ttk.Label(visualization, text='Waiting for microphone audio', anchor='center')
        self.spectrogram_label.pack(fill='both', expand=True)
        self.spectrogram_photo = None

        footer = ttk.Frame(outer)
        footer.pack(fill='x')
        ttk.Button(footer, text='Restart recorder app', command=self.restart_app).pack(side='left')
        ttk.Label(footer, text='Command window: 2.5 s after silence trimming. Save logs the shown prompt, not a speech transcript.',
                  wraplength=820).pack(side='left', padx=12)

    def _label_changed(self, _event=None):
        self._refresh_phrases()
        self.clear_take()

    def _refresh_phrases(self):
        label = self.label.get()
        phrases = RECORDER_VARIATIONS.get(label, (RECORDER_PHRASES[label],))
        self.phrase_box['values'] = phrases
        self.phrase.set(phrases[0])

    def refresh_devices(self):
        try:
            devices = sd.query_devices()
            self.input_devices = [
                (index, device['name'], float(device['default_samplerate']),
                 sd.query_hostapis(device['hostapi'])['name'])
                for index, device in enumerate(devices) if device['max_input_channels'] > 0
            ]
            self.output_devices = [
                (index, device['name'], float(device['default_samplerate']),
                 sd.query_hostapis(device['hostapi'])['name'])
                for index, device in enumerate(devices) if device['max_output_channels'] > 0
            ]
            # A Windows microphone can be exposed through several PortAudio host APIs.
            # Show the API explicitly so users can distinguish duplicate endpoints.
            self.input_box['values'] = [f'{name} — {api} [{index}]' for index, name, _, api in self.input_devices]
            self.output_box['values'] = [f'{name} — {api} [{index}]' for index, name, _, api in self.output_devices]
            self.input_choice.set(self.input_box['values'][0] if self.input_devices else '')
            self.output_choice.set(self.output_box['values'][0] if self.output_devices else '')
            self.status.set(f"Found {len(self.input_devices)} microphone inputs and {len(self.output_devices)} playback outputs.")
        except Exception as exc:
            self.status.set(f'Audio device scan failed: {exc}')

    @staticmethod
    def _selected_device(choice, devices):
        selected = choice.get()
        for index, name, rate, api in devices:
            if selected == f'{name} — {api} [{index}]':
                return index, rate, name, api
        raise ValueError('Select an available audio device, or click Refresh devices.')

    def start_recording(self):
        if self.recording:
            return
        try:
            device, rate, name, api = self._selected_device(self.input_choice, self.input_devices)
            # Try the selected Windows endpoint first, then its aliases under other
            # host APIs. Some drivers enumerate an MME alias that cannot be opened.
            aliases = [item for item in self.input_devices if item[1].casefold() == name.casefold()]
            aliases.sort(key=lambda item: (item[0] != device, item[3] not in ('Windows WASAPI', 'Windows DirectSound')))
            failures = []
            stream = None
            for candidate_device, _, candidate_rate, candidate_api in aliases:
                candidate_rate = int(candidate_rate)
                try:
                    sd.check_input_settings(device=candidate_device, samplerate=candidate_rate,
                                            channels=1, dtype='float32')
                    stream = sd.InputStream(device=candidate_device, samplerate=candidate_rate,
                                            channels=1, dtype='float32', callback=self._audio_callback)
                    stream.start()
                    device, rate, api = candidate_device, candidate_rate, candidate_api
                    break
                except Exception as exc:
                    failures.append(f'{candidate_api}: {exc}')
                    if stream is not None:
                        try:
                            stream.close()
                        except Exception:
                            pass
                        stream = None
            if stream is None:
                details = '\n'.join(failures)
                raise RuntimeError(
                    f'Windows listed {name}, but none of its audio backends could open it. '
                    'Close any app using the microphone (including browser recorder tabs), '
                    'check Windows Settings > Privacy & security > Microphone and the mic mute/USB connection, '
                    'then use Refresh devices. Tried:\n' + details
                )
            with self.chunk_lock:
                self.chunks = []
            self.recorded_audio = None
            self.record_started_at = time.monotonic()
            self.stream = stream
            self.recording = True
            self.status.set('🔴 Recording — speak now. The spectrogram updates live; click Stop when finished.')
            self.start_button.configure(state='disabled')
            self.stop_button.configure(state='normal')
            self.save_button.configure(state='disabled')
            self.play_button.configure(state='disabled')
        except Exception as exc:
            self.status.set(f'Microphone did not start: {exc}')
            self.duration.set('Duration: 0.00 s')
            messagebox.showerror('Could not start recording', str(exc), parent=self.window)

    def _audio_callback(self, indata, frames, time_info, status):
        del frames, time_info
        if status:
            self.audio_warning = str(status)
        with self.chunk_lock:
            self.chunks.append(indata.copy())

    def _current_audio(self):
        with self.chunk_lock:
            chunks = list(self.chunks)
        if not chunks:
            return None
        return np.concatenate(chunks, axis=0).reshape(-1)

    def _update_live_view(self):
        if self.recording:
            if self.audio_warning:
                self.status.set(f'Recording with audio warning: {self.audio_warning}')
                self.audio_warning = None
            current = self._current_audio()
            elapsed = time.monotonic() - self.record_started_at
            self.duration.set(f'Duration: {elapsed:.2f} s')
            if elapsed >= 20:
                self.stop_recording()
                self.status.set('Maximum capture time (20 seconds) reached; review the clip and save if appropriate.')
                self.window.after(200, self._update_live_view)
                return
            if current is not None and len(current) > 400:
                rate = int(self.stream.samplerate)
                current = current[-int(3.0 * rate):]
                if rate != SR:
                    divisor = np.gcd(rate, SR)
                    shown = resample_poly(current, SR // divisor, rate // divisor)
                    shown_rate = SR
                else:
                    shown, shown_rate = current, rate
                max_samples = int(3.0 * shown_rate)
                shown = shown[-max_samples:]
                self._render_audio_spectrogram(shown, shown_rate, 'LIVE MICROPHONE SPECTROGRAM')
        self.window.after(200, self._update_live_view)

    def _render_audio_spectrogram(self, wave, rate, caption):
        if not len(wave):
            return
        nperseg = min(400, len(wave))
        noverlap = min(320, nperseg - 1)
        frequencies, _, power = scipy_spectrogram(
            wave, fs=rate, nperseg=nperseg, noverlap=noverlap, mode='psd',
        )
        decibels = 10 * np.log10(np.maximum(power, 1e-12))
        intensity = np.clip((decibels + 90) / 70, 0, 1)
        stops = np.array([0.0, .25, .5, .75, 1.0])
        palette = np.array([
            [0, 0, 4], [40, 15, 93], [145, 35, 111], [234, 110, 46], [252, 253, 191],
        ], dtype=np.float32)
        rgb = np.stack([np.interp(intensity, stops, palette[:, channel]) for channel in range(3)], axis=-1)
        rgb = np.flipud(np.asarray(rgb, dtype=np.uint8))
        image = Image.fromarray(rgb, mode='RGB').resize((1000, 300), Image.Resampling.BILINEAR)
        self.spectrogram_photo = ImageTk.PhotoImage(image)
        self.spectrogram_label.configure(image=self.spectrogram_photo, text=caption, compound='top')

    def stop_recording(self):
        if not self.recording:
            return
        try:
            self.stream.stop()
            rate = int(self.stream.samplerate)
            self.stream.close()
            self.stream = None
            self.recording = False
            wave = self._current_audio()
            if wave is None or not len(wave):
                raise ValueError('No audio frames arrived. Check the selected microphone and try again.')
            self.recorded_audio = (rate, wave)
            self._render_audio_spectrogram(wave, rate, 'CAPTURED AUDIO SPECTROGRAM')
            self.status.set(f'Recording stopped — {len(wave) / rate:.2f} s captured. Review the spectrogram and play it, then save or clear.')
            self.duration.set(f'Duration: {len(wave) / rate:.2f} s')
            self.start_button.configure(state='normal')
            self.stop_button.configure(state='disabled')
            self.play_button.configure(state='normal')
            self.save_button.configure(state='normal')
        except Exception as exc:
            self.status.set(f'Recording stopped with no usable audio: {exc}')
            self.start_button.configure(state='normal')
            self.stop_button.configure(state='disabled')
            messagebox.showerror('Recording failed', str(exc), parent=self.window)

    def play_take(self):
        if self.recorded_audio is None:
            return
        try:
            device, output_rate, _name, _api = self._selected_device(self.output_choice, self.output_devices)
            rate, wave = self.recorded_audio
            sd.stop()
            if int(output_rate) != rate:
                divisor = np.gcd(int(output_rate), rate)
                wave = resample_poly(wave, int(output_rate) // divisor, rate // divisor)
            sd.play(wave, samplerate=int(output_rate), device=device, blocking=False)
            self.status.set('Playing the unsaved take through the selected speaker.')
        except Exception as exc:
            self.status.set(f'Playback failed: {exc}')

    def save_take(self):
        if self.recording:
            messagebox.showwarning('Stop first', 'Stop the recording before saving.', parent=self.window)
            return
        try:
            status, _ = save_recording_data(
                self.root, self.recorded_audio, self.speaker.get().strip(), self.condition.get(),
                self.label.get(), self.consent.get(), self.phrase.get(),
            )
            self.status.set(status)
            self.recorded_audio = None
            self.save_button.configure(state='disabled')
            self.start_button.configure(state='normal')
            self.duration.set('Duration: 0.00 s')
        except ValueError as exc:
            messagebox.showerror('Recording not saved', str(exc), parent=self.window)
            self.status.set(f'Not saved: {exc}')
        except Exception as exc:
            messagebox.showerror('Recording save failed', str(exc), parent=self.window)
            self.status.set(f'Save failed: {exc}')

    def clear_take(self):
        if self.recording:
            self.stop_recording()
        self.recorded_audio = None
        with self.chunk_lock:
            self.chunks = []
        self.duration.set('Duration: 0.00 s')
        self.status.set('Ready for another take — click Start recording.')
        self.play_button.configure(state='disabled')
        self.save_button.configure(state='disabled')
        self.spectrogram_photo = None
        self.spectrogram_label.configure(image='', text='Waiting for microphone audio')

    def export_excel(self):
        try:
            path = excel_export(self.root)
            self.status.set(f'Excel manifest refreshed: {path}')
            messagebox.showinfo('Manifest updated', f'Workbook saved to:\n{path}', parent=self.window)
        except Exception as exc:
            messagebox.showerror('Excel export failed', f'CSV and WAV recordings remain safe.\n\n{exc}', parent=self.window)

    def restart_app(self):
        if self.recording:
            messagebox.showwarning('Stop first', 'Stop or save the current take before restarting.', parent=self.window)
            return
        self.close(restart=True)

    def close(self, restart=False):
        if self.recording and not messagebox.askyesno('Discard current take?', 'Stop and discard the unsaved recording?', parent=self.window):
            return
        try:
            if self.stream is not None:
                self.stream.stop()
                self.stream.close()
                self.stream = None
            sd.stop()
            self.window.destroy()
        finally:
            if restart:
                os.execv(sys.executable, [sys.executable, *sys.argv])

    def run(self):
        self.window.mainloop()


def main():
    try:
        app = DesktopRecorder(project_root())
        app.run()
    except Exception as exc:
        try:
            messagebox.showerror(APP_TITLE, str(exc))
        except Exception:
            raise


if __name__ == '__main__':
    main()
